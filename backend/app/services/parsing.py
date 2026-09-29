"""Resume text extraction and lightweight structured parsing.

Text extraction uses pdfplumber/PyPDF2 for PDF and python-docx for DOCX.
Structured parsing is intentionally dependency-free (regex + heuristics) so it
stays fast and deterministic.
"""

import io
import re

import pdfplumber
from docx import Document
from PyPDF2 import PdfReader

SECTION_ALIASES = {
    "summary": ["summary", "objective", "profile", "about me", "professional summary"],
    "experience": [
        "experience",
        "work experience",
        "professional experience",
        "employment",
        "work history",
    ],
    "education": ["education", "academics", "academic background", "qualifications"],
    "skills": ["skills", "technical skills", "core competencies", "technologies"],
    "projects": ["projects", "personal projects", "academic projects"],
    "certifications": ["certifications", "certificates", "licenses"],
}

SKILL_VOCAB = [
    "python", "java", "javascript", "typescript", "c++", "c#", "go", "rust", "php",
    "ruby", "kotlin", "swift", "scala", "r", "matlab", "sql", "nosql", "mongodb",
    "postgresql", "mysql", "sqlite", "redis", "elasticsearch", "cassandra",
    "react", "next.js", "nextjs", "vue", "angular", "svelte", "node.js", "express",
    "fastapi", "django", "flask", "spring", "laravel", "tailwind", "html", "css",
    "rest", "graphql", "grpc", "microservices", "docker", "kubernetes", "terraform",
    "ansible", "jenkins", "github actions", "ci/cd", "aws", "azure", "gcp",
    "machine learning", "deep learning", "nlp", "tensorflow", "pytorch", "scikit-learn",
    "pandas", "numpy", "opencv", "llm", "langchain", "openai", "groq", "spacy",
    "git", "linux", "bash", "agile", "scrum", "jira", "figma", "prisma", "jwt",
    "oauth", "unit testing", "pytest", "jest", "selenium", "cypress", "devops",
]

EMAIL_RE = re.compile(r"[\w.+-]+@[\w-]+\.[\w.-]+")
PHONE_RE = re.compile(r"(?:(?:\+?\d{1,3}[\s-]?)?(?:\(?\d{3}\)?[\s.-]?)\d{3}[\s.-]?\d{4})")
LINKEDIN_RE = re.compile(r"(?:https?://)?(?:www\.)?linkedin\.com/[\w/\-]+", re.I)
GITHUB_RE = re.compile(r"(?:https?://)?(?:www\.)?github\.com/[\w/\-]+", re.I)
URL_RE = re.compile(r"https?://[\w./\-?=&%#]+", re.I)


def extract_text(file_bytes: bytes, filename: str) -> str:
    lower = filename.lower()
    if lower.endswith(".pdf"):
        return _extract_pdf(file_bytes)
    if lower.endswith(".docx"):
        return _extract_docx(file_bytes)
    if lower.endswith(".txt"):
        return file_bytes.decode("utf-8", errors="ignore")
    raise ValueError("Unsupported file type. Upload a PDF, DOCX or TXT file.")


def _extract_pdf(file_bytes: bytes) -> str:
    text_parts: list[str] = []
    try:
        with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
            for page in pdf.pages:
                text_parts.append(page.extract_text() or "")
    except Exception:
        text_parts = []

    text = "\n".join(text_parts).strip()
    if text:
        return text

    # Fallback for PDFs pdfplumber cannot read (encrypted / unusual encoding).
    reader = PdfReader(io.BytesIO(file_bytes))
    return "\n".join((page.extract_text() or "") for page in reader.pages).strip()


def _extract_docx(file_bytes: bytes) -> str:
    document = Document(io.BytesIO(file_bytes))
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(parts).strip()


def _detect_sections(lines: list[str]) -> dict[str, str]:
    """Group lines into canonical resume sections using header aliases."""
    alias_to_section = {
        alias: section
        for section, aliases in SECTION_ALIASES.items()
        for alias in aliases
    }
    header_re = re.compile(r"^[A-Za-z][A-Za-z &/]{2,40}$")

    sections: dict[str, list[str]] = {}
    current = "header"
    sections[current] = []
    for line in lines:
        stripped = line.strip()
        if not stripped:
            continue
        normalized = stripped.lower().strip(":•- ")
        if header_re.match(stripped) and normalized in alias_to_section:
            current = alias_to_section[normalized]
            sections.setdefault(current, [])
            continue
        sections.setdefault(current, []).append(stripped)

    return {name: "\n".join(content) for name, content in sections.items() if content}


def parse_resume(raw_text: str) -> dict:
    lines = [ln for ln in raw_text.splitlines()]
    non_empty = [ln.strip() for ln in lines if ln.strip()]
    sections = _detect_sections(lines)
    lower_text = raw_text.lower()

    email_match = EMAIL_RE.search(raw_text)
    phone_match = PHONE_RE.search(raw_text)
    linkedin_match = LINKEDIN_RE.search(raw_text)
    github_match = GITHUB_RE.search(raw_text)
    links = list({u for u in URL_RE.findall(raw_text)})

    name = _guess_name(non_empty)
    skills = sorted(
        {skill for skill in SKILL_VOCAB if re.search(rf"(?<!\w){re.escape(skill)}(?!\w)", lower_text)}
    )

    return {
        "name": name,
        "email": email_match.group(0) if email_match else None,
        "phone": phone_match.group(0).strip() if phone_match else None,
        "linkedin": linkedin_match.group(0) if linkedin_match else None,
        "github": github_match.group(0) if github_match else None,
        "links": links[:5],
        "sections": list(sections.keys()),
        "sectionContent": sections,
        "skills": skills,
        "experience": _extract_bullets(sections.get("experience", "")),
        "education": _extract_bullets(sections.get("education", "")),
        "projects": _extract_bullets(sections.get("projects", "")),
        "wordCount": len(raw_text.split()),
        "lineCount": len(non_empty),
    }


def _guess_name(non_empty_lines: list[str]) -> str | None:
    """The name is normally the first short, title-cased line without digits/@."""
    for line in non_empty_lines[:5]:
        if "@" in line or any(ch.isdigit() for ch in line):
            continue
        if 2 <= len(line.split()) <= 4 and len(line) <= 50:
            if line.replace(".", "").replace(" ", "").isalpha() or line.istitle():
                return line
    return None


def _extract_bullets(section_text: str) -> list[str]:
    if not section_text:
        return []
    bullets = []
    for line in section_text.splitlines():
        cleaned = line.strip(" \t•-*")
        if len(cleaned) > 3:
            bullets.append(cleaned)
    return bullets[:15]
